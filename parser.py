from dataclasses import dataclass
from typing import List
from openpyxl import load_workbook
from datetime import date, datetime
import requests
import tempfile
import os
from models import ReportItem
PROGRAM_MAP = {
    "Cab Retreat": "6100",
    "Cab Food/Prizes": "6653",
    "Laundry": "8206",
    "Initiative": "0200",
    "Storage Assistance": "0330",
    "Socioeconomic Inclusivity Fund": "1350",
    "RAs Budget": "6658",
    "Awards": "6333",
    "Spirit": "6137",
    "Sports": "2400",
    "Beer Bike": "6133",
    "Merch": "5080",
    "Community Service": "0331",
    "Permanent Improvements": "6531",
    "Staff Appreciation": "8601",
    "Associates": "6650",
    "Alumni": "6951",
    "Academic Mentors": "1110",
    "PAAs": "6655",
    "Diversity": "6600",
    "Senior Class": "6136",
    "Junior Class": "6601",
    "Sophomore Class": "6602",
    "Freshmen Class": "6134",
    "Off Campus": "6135",
    "Socials": "4600",
    "BGHS": "6603",
    "Chief Justice": "8202",
    "President": "6656",
    "STRIVE": "6604",
    "RHA": "8604",
    "Baker Shake": "3150",
    "Baker Black Caucus": "6605",
    "O-Week": "6000",
}

def export_google_sheet_to_workbook(sheet_url: str):
    spreadsheet_id = sheet_url.split("/d/")[1].split("/")[0]
    export_url = (
        f"https://docs.google.com/spreadsheets/d/"
        f"{spreadsheet_id}/export?format=xlsx"
    )

    response = requests.get(export_url)
    response.raise_for_status()

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.write(response.content)
    tmp.close()

    wb = load_workbook(tmp.name)
    os.unlink(tmp.name)
    return wb


def parse_date_flexible(date_str: str) -> date:
    """
    Accepts M/D/YYYY, M/D/YY, or M/D (assumes current year).
    """
    parts = date_str.strip().split("/")
    month = int(parts[0])
    day = int(parts[1])
    if len(parts) == 3 and parts[2]:
        year = int(parts[2])
        if year < 100:
            year += 2000
    else:
        year = date.today().year
    return date(year, month, day)

def parse_receipts(cell_value) -> List[str]:
    if not cell_value:
        return []

    return [
        part.strip()
        for part in str(cell_value).split(",")
        if part.strip()
    ]

def extract_date(cell_value) -> date:
    if isinstance(cell_value, datetime):
        return cell_value.date()

    if isinstance(cell_value, date):
        return cell_value

    s = str(cell_value).strip()
    if "/" in s:
        return parse_date_flexible(s)
    return datetime.fromisoformat(s.split(" ")[0]).date()

def parse_price(cell_value) -> float:
    if not cell_value:
        return 0.0

    if isinstance(cell_value, (int, float)):
        return float(cell_value)

    price_str = str(cell_value).strip().replace("$", "").replace(",", "")

    try:
        return float(price_str)
    except ValueError:
        return 0.0

def parse_purchases(
    spreadsheet_link: str,
    cardholder_name: str,
    start_date: str
) -> List[ReportItem]:

    wb = export_google_sheet_to_workbook(spreadsheet_link)

    print(f"[parser] Available sheets: {wb.sheetnames}")
    ws = wb["Purchases"]

    start_dt = parse_date_flexible(start_date)
    print(f"[parser] Parsed start_date: {start_dt}")
    print(f"[parser] Filtering for cardholder: '{cardholder_name}'")
    print(f"[parser] Total rows in sheet (including header): {ws.max_row}")

    report_items: List[ReportItem] = []
    rows_scanned = 0
    rows_skipped_no_timestamp = 0
    rows_skipped_before_date = 0
    rows_skipped_wrong_cardholder = 0
    unique_cardholders_seen: set = set()

    for row in range(ws.max_row, 1, -1):
        A_timestamp = ws[f"A{row}"].value
        G_pcard = ws[f"G{row}"].value

        if not A_timestamp:
            rows_skipped_no_timestamp += 1
            continue

        rows_scanned += 1

        row_date = extract_date(A_timestamp)

        if row_date < start_dt:
            rows_skipped_before_date += 1
            print(f"[parser] Row {row}: date {row_date} < start_date {start_dt} — stopping scan")
            break

        if G_pcard:
            unique_cardholders_seen.add(str(G_pcard))

        if G_pcard != cardholder_name:
            rows_skipped_wrong_cardholder += 1
            continue

        B_name = ws[f"B{row}"].value
        C_items = ws[f"C{row}"].value
        D_vendor = ws[f"D{row}"].value
        E_event = ws[f"E{row}"].value
        F_price = ws[f"F{row}"].value
        H_budget = ws[f"H{row}"].value
        I_endowment = ws[f"I{row}"].value
        J_flyer = ws[f"J{row}"].value
        K_receipts = ws[f"K{row}"].value
        L_needs_affidavit = ws[f"L{row}"].value

        description_parts = [
            str(B_name),
            "Baker College",
            row_date.strftime("%m/%d/%Y"),
            str(D_vendor),
            str(E_event),
            str(C_items),
            str(H_budget),
        ]

        if H_budget == "Other" and I_endowment:
            description_parts.append(str(I_endowment))

        description = " | ".join(description_parts)
        activity = PROGRAM_MAP.get(H_budget, "")


        report_items.append(
            ReportItem(
                description=description,
                activity=activity,
                date=row_date,
                price=parse_price(F_price),
                vendor=str(D_vendor) if D_vendor else "",
                receipts=parse_receipts(K_receipts),
                flyer=str(J_flyer) if J_flyer else "",
                needsAffidavit=(L_needs_affidavit == "No")
            )
        )

    print(f"[parser] Rows scanned (with timestamp): {rows_scanned}")
    print(f"[parser] Rows skipped (no timestamp): {rows_skipped_no_timestamp}")
    print(f"[parser] Rows skipped (before start_date): {rows_skipped_before_date}")
    print(f"[parser] Rows skipped (wrong cardholder): {rows_skipped_wrong_cardholder}")
    print(f"[parser] Unique cardholders seen in range: {unique_cardholders_seen}")
    print(f"[parser] Items matched for '{cardholder_name}': {len(report_items)}")

    return report_items
